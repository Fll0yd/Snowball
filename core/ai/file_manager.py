import os
import hashlib
import threading
import time
from queue import PriorityQueue
import tensorflow as tf
from tensorflow.keras.models import load_model
from tensorflow.keras.preprocessing.image import load_img, img_to_array
from PyPDF2 import PdfReader
import docx
from PIL import Image
import pytesseract
from watchdog.observers import Observer
from watchdog.events import FileSystemEventHandler
from datetime import datetime
import csv
import openpyxl

class FileEventHandler(FileSystemEventHandler):
    """Custom event handler for file system changes."""

    def __init__(self, logger, priority_queue):
        self.logger = logger
        self.priority_queue = priority_queue

    def on_created(self, event):
        if not event.is_directory:
            self.logger.log_task(f"File created: {event.src_path}", "Created")
            self.priority_queue.put((1, event.src_path))  # Low priority for creation

    def on_modified(self, event):
        if not event.is_directory:
            self.logger.log_task(f"File modified: {event.src_path}", "Modified")
            self.priority_queue.put((0, event.src_path))  # High priority for modification

class FileManager:
    def __init__(self, logger, memory, scan_dir='S:/', allowed_extensions=None, scan_interval=10):
        self.logger = logger
        self.memory = memory
        self.scan_dir = scan_dir
        self.priority_queue = PriorityQueue()
        self.processed_hashes = set()
        self.hash_lock = threading.Lock()
        self.observer = Observer()
        self.allowed_extensions = allowed_extensions or {
            '.txt', '.py', '.json', '.log', '.pdf', '.docx', '.jpg', '.jpeg', '.png'
        }
        self.scan_interval = scan_interval
        self.stop_event = threading.Event()

        # Initialize cooldown mechanism
        self.last_processed_files = {}  # Dictionary to track cooldowns
        self.cooldown_period = 2  # Cooldown period in seconds

        # Load pre-trained models
        self.image_model = self.load_image_model()
        self.text_model = self.load_text_model()
        
    def load_image_model(self):
        """Load a pre-trained image classification model."""
        try:
            model = load_model('S:/Snowball/models/image_classification_model.h5')
            self.logger.log_task("Image classification model loaded successfully.", "Loaded")
            return model
        except Exception as e:
            self.logger.log_error(f"Error loading image classification model: {e}")
            return None

    def load_text_model(self):
        """Load a pre-trained text classification model."""
        try:
            model = load_model('S:/Snowball/models/text_classification_model.h5')
            self.logger.log_task("Text classification model loaded successfully.", "Loaded")
            return model
        except Exception as e:
            self.logger.log_error(f"Error loading text classification model: {e}")
            return None
        
    def validate_models(self):
        if not os.path.exists('S:/Snowball/models/text_classification_model.h5'):
            self.logger.log_error("Text classification model is missing.")
        if not os.path.exists('S:/Snowball/models/image_classification_model.h5'):
            self.logger.log_error("Image classification model is missing.")

    def hash_file(self, file_path):
        hasher = hashlib.md5()
        try:
            with open(file_path, 'rb') as f:
                hasher.update(f.read())
            return hasher.hexdigest()
        except Exception as e:
            self.logger.log_error(f"Error hashing file: {e}")
            return None

    def scan_and_index_drive(self, batch_size=100):
        """Scan and index all files in batches."""
        try:
            self.logger.log_event("Starting drive scan and index.")
            files_to_process = []
            for root, _, files in os.walk(self.scan_dir):
                for file_name in files:
                    if os.path.splitext(file_name)[1].lower() not in self.allowed_extensions:
                        continue
                    file_path = os.path.join(root, file_name)
                    files_to_process.append(file_path)
                    if len(files_to_process) >= batch_size:
                        self._process_batch(files_to_process)
                        files_to_process.clear()
            if files_to_process:
                self._process_batch(files_to_process)
            self.logger.log_event("Drive scan and indexing completed.")
        except Exception as e:
            self.logger.log_error(f"Error scanning and indexing drive: {e}")
    
    def analyze_file(self, file_path):
        """Analyze a file based on its type."""
        ext = os.path.splitext(file_path)[1].lower()

        if ext == '.pdf':
            self.analyze_pdf(file_path)
        elif ext == '.docx':
            self.analyze_docx(file_path)
        elif ext == '.txt':
            self.analyze_text(file_path)
        elif ext in ['.jpg', '.jpeg', '.png']:
            self.analyze_image(file_path)
        elif ext == '.csv':
            self.analyze_csv(file_path)
        elif ext == '.xlsx':
            self.analyze_excel(file_path)
        else:
            self.logger.log_warning(f"Unsupported file type skipped: {file_path}")

    def analyze_pdf(self, file_path):
        """Extract and analyze text from a PDF file."""
        if not self.text_model:
            self.logger.log_warning(f"Text model unavailable. Skipping PDF analysis for {file_path}.")
            return

        try:
            reader = PdfReader(file_path)
            text = ''.join(page.extract_text() for page in reader.pages)

            prediction = self.text_model.predict([text])
            document_type = "Important" if prediction[0] > 0.5 else "General"

            self.logger.log_task(f"Analyzed PDF file: {file_path} (Type: {document_type})", "Analyzed")
            self.memory.store_file_analysis(file_path, text)
        except Exception as e:
            self.logger.log_error(f"Error analyzing PDF file: {file_path}, {e}")

    def analyze_docx(self, file_path):
        """Extract and analyze text from a DOCX file."""
        if not self.text_model:
            self.logger.log_warning(f"Text model unavailable. Skipping DOCX analysis for {file_path}.")
            return

        try:
            doc = docx.Document(file_path)
            text = ' '.join(paragraph.text for paragraph in doc.paragraphs)

            prediction = self.text_model.predict([text])
            document_type = "Important" if prediction[0] > 0.5 else "General"

            self.logger.log_task(f"Analyzed DOCX file: {file_path} (Type: {document_type})", "Analyzed")
            self.memory.store_file_analysis(file_path, text)
        except Exception as e:
            self.logger.log_error(f"Error analyzing DOCX file: {file_path}, {e}")

    def analyze_text(self, file_path):
        """Extract and analyze text from a TXT file."""
        if not self.text_model:
            self.logger.log_warning(f"Text model unavailable. Skipping TXT analysis for {file_path}.")
            return

        try:
            with open(file_path, 'r', encoding='utf-8') as f:
                text = f.read()

            prediction = self.text_model.predict([text])
            document_type = "Important" if prediction[0] > 0.5 else "General"

            self.logger.log_task(f"Analyzed TXT file: {file_path} (Type: {document_type})", "Analyzed")
            self.memory.store_file_analysis(file_path, text)
        except Exception as e:
            self.logger.log_error(f"Error analyzing TXT file: {file_path}, {e}")

    def analyze_image(self, file_path):
        """Perform OCR and image classification on an image file."""
        if not self.image_model:
            self.logger.log_warning(f"Image model unavailable. Skipping image analysis for {file_path}.")
            return

        try:
            image = Image.open(file_path)
            text = pytesseract.image_to_string(image)

            img = load_img(file_path, target_size=(224, 224))
            img_array = img_to_array(img) / 255.0
            img_array = img_array.reshape((1, *img_array.shape))
            prediction = self.image_model.predict(img_array)
            classification = "Relevant" if prediction[0] > 0.5 else "Irrelevant"

            self.logger.log_task(f"Analyzed Image file: {file_path} (Classification: {classification})", "Analyzed")
            self.memory.store_file_analysis(file_path, text)
        except Exception as e:
            self.logger.log_error(f"Error analyzing Image file: {file_path}, {e}")

    def analyze_csv(self, file_path):
        """Analyze a CSV file."""
        try:
            with open(file_path, 'r') as f:
                reader = csv.reader(f)
                rows = list(reader)
                self.logger.log_task(f"Analyzed CSV file: {file_path} with {len(rows)} rows.", "Analyzed")
                self.memory.store_file_analysis(file_path, str(rows))
        except Exception as e:
            self.logger.log_error(f"Error analyzing CSV file: {file_path}, {e}")

    def analyze_excel(self, file_path):
        """Analyze an Excel file."""
        try:
            workbook = openpyxl.load_workbook(file_path)
            text = []
            for sheet in workbook:
                for row in sheet.iter_rows(values_only=True):
                    text.append(' '.join(map(str, row)))

            self.logger.log_task(f"Analyzed Excel file: {file_path} with {len(text)} rows.", "Analyzed")
            self.memory.store_file_analysis(file_path, '\n'.join(text))
        except Exception as e:
            self.logger.log_error(f"Error analyzing Excel file: {file_path}, {e}")

    def _process_batch(self, files):
        try:
            log_messages = []
            for file_path in files:
                last_modified = datetime.fromtimestamp(os.path.getmtime(file_path)).isoformat()
                self.memory.store_file_metadata(os.path.basename(file_path), file_path, last_modified)
                log_messages.append(f"Indexed file: {file_path}")
            if log_messages:
                self.logger.log_event(f"Batch processed {len(files)} files:\n" + "\n".join(log_messages))
        except Exception as e:
            self.logger.log_error(f"Error processing batch: {e}")

    def should_process_file(self, file_path):
        current_time = time.time()
        if file_path in self.last_processed_files:
            if current_time - self.last_processed_files[file_path] < self.cooldown_period:
                return False
        self.last_processed_files[file_path] = current_time
        return True

    def process_file(self, file_path):
        """Analyze a file based on its type."""
        try:
            file_hash = self.hash_file(file_path)
            with self.hash_lock:
                if not file_hash or file_hash in self.processed_hashes:
                    self.logger.log_task(f"File already processed: {file_path}", "Skipped")
                    return

            ext = os.path.splitext(file_path)[1].lower()
            if ext == '.pdf':
                self.analyze_pdf(file_path)
            elif ext == '.docx':
                self.analyze_docx(file_path)
            elif ext == '.txt':
                self.analyze_text(file_path)
            elif ext in ['.jpg', '.jpeg', '.png']:
                self.analyze_image(file_path)
            elif ext == '.csv':
                self.analyze_csv(file_path)
            elif ext == '.xlsx':
                self.analyze_excel(file_path)
            else:
                self.logger.log_warning(f"Unsupported file type skipped: {file_path} (Extension: {ext})")

            with self.hash_lock:
                self.processed_hashes.add(file_hash)
        except Exception as e:
            self.logger.log_error(f"Error processing file: {file_path}, {e}")

    def start_monitoring(self):
        if not os.path.exists(self.scan_dir):
            self.logger.log_error(f"Directory {self.scan_dir} does not exist.")
            return
        event_handler = FileEventHandler(self.logger, self.priority_queue)
        self.observer.schedule(event_handler, self.scan_dir, recursive=True)
        self.observer.start()
        threading.Thread(target=self._monitor_files, daemon=True).start()

    def _monitor_files(self):
        """Continuously monitor files."""
        try:
            while not self.stop_event.is_set():
                try:
                    _, file_path = self.priority_queue.get(timeout=self.scan_interval)
                    if os.path.exists(file_path):
                        self.process_file(file_path)
                except queue.Empty:
                    continue
        except Exception as e:
            self.logger.log_error(f"Error in monitoring loop: {e}")

    def stop_monitoring(self):
        self.stop_event.set()
        self.logger.log_event("Draining remaining files before shutdown.")
        while not self.priority_queue.empty():
            _, file_path = self.priority_queue.get()
            if os.path.exists(file_path):
                self.process_file(file_path)
        self.observer.stop()
        self.observer.join()
        self.logger.log_event("Monitoring stopped.")

